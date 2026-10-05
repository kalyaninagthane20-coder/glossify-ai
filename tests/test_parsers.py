from io import BytesIO

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.ingestion.parsers import UnsupportedFormatError, parse_document


def test_text_file_keeps_headings_bullets_and_paragraphs():
    text = b"# Title\n\nFirst line\nsecond line\n\n- item one\n* item two\n"
    blocks = parse_document("notes.md", text)
    assert [(b.block_type, b.text) for b in blocks] == [
        ("heading", "Title"),
        ("paragraph", "First line second line"),
        ("list_item", "item one"),
        ("list_item", "item two"),
    ]
    assert blocks[0].heading_level == 1


def test_docx_keeps_headings_lists_and_table_positions():
    doc = DocxDocument()
    doc.add_heading("Introduction", level=2)
    doc.add_paragraph("Hello world.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Qty"
    table.cell(1, 0).text = "Pen"
    table.cell(1, 1).text = "3"
    doc.add_paragraph("Shopping list", style="List Bullet")
    buffer = BytesIO()
    doc.save(buffer)

    blocks = parse_document("sample.docx", buffer.getvalue())

    assert blocks[0].block_type == "heading" and blocks[0].heading_level == 2
    assert blocks[1].block_type == "paragraph"
    cells = [b for b in blocks if b.block_type == "table_cell"]
    assert [(c.row_index, c.col_index, c.text) for c in cells] == [
        (0, 0, "Name"), (0, 1, "Qty"), (1, 0, "Pen"), (1, 1, "3")]
    assert all(c.table_index == 0 for c in cells)
    assert blocks[-1].block_type == "list_item"


def test_pdf_text_is_extracted_into_paragraphs():
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, "The first paragraph describes the invoice process.")
    pdf.ln(4)
    pdf.multi_cell(0, 8, "The second paragraph covers payment terms.")
    blocks = parse_document("report.pdf", bytes(pdf.output()))

    joined = " ".join(b.text for b in blocks)
    assert "first paragraph" in joined
    assert "payment terms" in joined
    assert all(b.page == 1 for b in blocks)


def test_unsupported_extension_is_rejected():
    with pytest.raises(UnsupportedFormatError):
        parse_document("image.png", b"\x89PNG")
