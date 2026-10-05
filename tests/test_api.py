import io

from docx import Document as DocxDocument


def upload_txt(client, text=b"# Contract\n\nPlease send the invoice 42.\n\nThank you.",
               name="contract.txt", source="en", target="fr"):
    return client.post(
        "/api/documents",
        files={"file": (name, text, "text/plain")},
        data={"source_lang": source, "target_lang": target},
    )


def test_health_and_index_page(client):
    assert client.get("/api/health").json() == {"status": "ok"}
    assert client.get("/").status_code == 200


def test_full_workflow_upload_translate_review_export(client):
    upload = upload_txt(client)
    assert upload.status_code == 201
    document = upload.json()
    assert document["segment_count"] == 3
    assert [s["block_type"] for s in document["segments"]] == ["heading", "paragraph", "paragraph"]

    report = client.post(f"/api/documents/{document['id']}/translate").json()
    assert report["translated"] == 3
    assert report["by_method"] == {"upper": 3}

    # export is blocked until every segment is approved
    blocked = client.get(f"/api/documents/{document['id']}/export?format=txt")
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["unapproved_segments"] == 3

    for segment in client.get(f"/api/documents/{document['id']}").json()["segments"]:
        approved = client.post(f"/api/segments/{segment['id']}/approve")
        assert approved.status_code == 200
        assert approved.json()["status"] == "approved"

    exported = client.get(f"/api/documents/{document['id']}/export?format=txt")
    assert exported.status_code == 200
    assert "PLEASE SEND THE INVOICE 42." in exported.text
    assert exported.text.startswith("# CONTRACT")


def test_docx_export_is_a_word_file_with_the_same_structure(client):
    document = upload_txt(client).json()
    client.post(f"/api/documents/{document['id']}/translate")
    for segment in client.get(f"/api/documents/{document['id']}").json()["segments"]:
        client.post(f"/api/segments/{segment['id']}/approve")

    response = client.get(f"/api/documents/{document['id']}/export?format=docx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"  # .docx is a zip container
    assert "contract_fr.docx" in response.headers["content-disposition"]

    exported = DocxDocument(io.BytesIO(response.content))
    assert exported.paragraphs[0].text == "CONTRACT"
    assert exported.paragraphs[0].style.name == "Heading 1"


def test_approved_pairs_feed_the_translation_memory(client):
    document = upload_txt(client, text=b"Pay the invoice.", name="pay.txt").json()
    client.post(f"/api/documents/{document['id']}/translate")
    segment_id = document["segments"][0]["id"]
    client.post(f"/api/segments/{segment_id}/approve")

    matches = client.get("/api/tm/search", params={"q": "pay the invoice", "source_lang": "en",
                                                   "target_lang": "fr"}).json()
    assert matches[0]["target_text"] == "PAY THE INVOICE."


def test_editing_a_segment_reruns_qa_and_sets_in_review(client):
    client.post("/api/glossary", json={"source_lang": "en", "target_lang": "fr",
                                       "source_term": "invoice", "target_term": "facture"})
    document = upload_txt(client, text=b"Send the invoice.", name="a.txt").json()
    client.post(f"/api/documents/{document['id']}/translate")
    segment_id = document["segments"][0]["id"]

    translated = client.get(f"/api/documents/{document['id']}").json()["segments"][0]
    assert translated["status"] == "needs_review"  # "SEND THE INVOICE." lacks "facture"

    edited = client.patch(f"/api/segments/{segment_id}",
                          json={"target_text": "Envoyez la facture.", "comment": "terminology"}).json()
    assert edited["status"] == "in_review"
    assert edited["qa_issues"] == []
    assert edited["method"] == "manual"


def test_reject_records_the_comment(client):
    document = upload_txt(client, text=b"Hello there.", name="r.txt").json()
    segment_id = document["segments"][0]["id"]
    rejected = client.post(f"/api/segments/{segment_id}/reject", json={"comment": "wrong register"}).json()
    assert rejected["status"] == "rejected"
    assert rejected["reviewer_comment"] == "wrong register"


def test_glossary_crud(client):
    created = client.post("/api/glossary", json={"source_lang": "EN", "target_lang": "fr",
                                                 "source_term": "invoice", "target_term": "facture"})
    assert created.status_code == 201
    term_id = created.json()["id"]
    assert created.json()["source_lang"] == "en"

    listed = client.get("/api/glossary", params={"source_lang": "en", "target_lang": "fr"}).json()
    assert [t["source_term"] for t in listed] == ["invoice"]

    assert client.delete(f"/api/glossary/{term_id}").status_code == 204
    assert client.delete(f"/api/glossary/{term_id}").status_code == 404


def test_unsupported_file_type_is_rejected(client):
    response = client.post("/api/documents", files={"file": ("photo.png", b"\x89PNG", "image/png")},
                           data={"source_lang": "en", "target_lang": "fr"})
    assert response.status_code == 415


def test_unknown_ids_return_404(client):
    assert client.get("/api/documents/999").status_code == 404
    assert client.patch("/api/segments/999", json={"target_text": "x"}).status_code == 404
    assert client.post("/api/segments/999/approve").status_code == 404
