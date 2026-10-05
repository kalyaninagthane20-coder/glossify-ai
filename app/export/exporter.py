"""Export a document back to .txt or .docx, keeping the original structure.

Headings keep their level, list items stay as list items, and table cells are
rebuilt as a table with the same number of rows and columns.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import Iterator, Sequence

from docx import Document as DocxDocument

from app.models import Segment


def _text(segment: Segment) -> str:
    return segment.target_text or segment.source_text


def _units(segments: Sequence[Segment]) -> Iterator[tuple[str, object]]:
    """Yield ('block', segment) or ('table', {row: {col: text}}) in document order."""
    tables: dict[int, dict[int, dict[int, str]]] = defaultdict(lambda: defaultdict(dict))
    for seg in segments:
        if seg.block_type == "table_cell":
            tables[seg.table_index][seg.row_index][seg.col_index] = _text(seg)

    emitted: set[int] = set()
    for seg in segments:
        if seg.block_type != "table_cell":
            yield "block", seg
        elif seg.table_index not in emitted:
            emitted.add(seg.table_index)
            yield "table", tables[seg.table_index]


def export_text(segments: Sequence[Segment]) -> str:
    lines: list[str] = []
    for kind, payload in _units(segments):
        if kind == "table":
            for row in sorted(payload):
                cells = payload[row]
                lines.append(" | ".join(cells[col] for col in sorted(cells)))
            lines.append("")
        elif payload.block_type == "heading":
            lines.append(f"{'#' * (payload.heading_level or 1)} {_text(payload)}")
            lines.append("")
        elif payload.block_type == "list_item":
            lines.append(f"- {_text(payload)}")
        else:
            lines.append(_text(payload))
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def export_docx(segments: Sequence[Segment]) -> bytes:
    document = DocxDocument()
    for kind, payload in _units(segments):
        if kind == "table":
            n_rows = max(payload) + 1
            n_cols = max(max(cells) for cells in payload.values()) + 1
            table = document.add_table(rows=n_rows, cols=n_cols)
            table.style = "Table Grid"
            for row, cells in payload.items():
                for col, text in cells.items():
                    table.cell(row, col).text = text
        elif payload.block_type == "heading":
            document.add_heading(_text(payload), level=min(max(payload.heading_level or 1, 1), 9))
        elif payload.block_type == "list_item":
            document.add_paragraph(_text(payload), style="List Bullet")
        else:
            document.add_paragraph(_text(payload))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
